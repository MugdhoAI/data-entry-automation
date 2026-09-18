"""Input loading, cleaning, validation, and export."""

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .validators import clean_value, normalize_email, normalize_phone, validate_record

REQUIRED_COLUMNS = ("name", "email", "phone")
OPTIONAL_COLUMNS = ("company",)


class ProcessingResult:
    """Result of processing one input file."""

    def __init__(
        self,
        records: list[dict[str, str]],
        errors: list[dict[str, str | int]],
        total_rows: int,
    ) -> None:
        self.records = records
        self.errors = errors
        self.total_rows = total_rows


def _normalize_headers(headers: list[object]) -> list[str]:
    return [clean_value(value).lower().replace(" ", "_") for value in headers]


def _normalize_record(row: dict[str, object]) -> dict[str, str]:
    return {
        "name": clean_value(row.get("name")),
        "email": normalize_email(row.get("email")),
        "phone": normalize_phone(row.get("phone")),
        "company": clean_value(row.get("company")),
    }


def _validate_headers(headers: list[str]) -> None:
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")


def _read_csv(path: Path) -> list[dict[str, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = _normalize_headers(reader.fieldnames or [])
        _validate_headers(headers)
        return [
            dict(zip(headers, row.values(), strict=False))
            for row in reader
        ]


def _read_xlsx(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = _normalize_headers(list(next(rows, [])))
        _validate_headers(headers)
        return [
            dict(zip(headers, row, strict=False))
            for row in rows
            if any(value is not None and str(value).strip() for value in row)
        ]
    finally:
        workbook.close()


def read_input(path: Path) -> list[dict[str, object]]:
    """Read CSV or XLSX input based on its file extension."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    raise ValueError("Unsupported input format. Use .csv or .xlsx.")


def process_file(path: Path) -> ProcessingResult:
    """Clean and validate every row, separating valid and invalid records."""
    rows = read_input(path)
    valid: list[dict[str, str]] = []
    errors: list[dict[str, str | int]] = []

    for row_number, row in enumerate(rows, start=2):
        record = _normalize_record(row)
        row_errors = validate_record(record, row_number)
        if row_errors:
            errors.extend(
                {"row": error.row, "field": error.field, "message": error.message}
                for error in row_errors
            )
        else:
            valid.append(record)

    return ProcessingResult(valid, errors, len(rows))


def write_csv(path: Path, records: list[dict[str, str]]) -> None:
    """Write cleaned records to CSV."""
    fieldnames = [*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)


def write_xlsx(path: Path, records: list[dict[str, str]]) -> None:
    """Write cleaned records to Excel."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cleaned Data"
    fieldnames = [*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS]
    sheet.append(fieldnames)
    for record in records:
        sheet.append([record.get(field, "") for field in fieldnames])
    workbook.save(path)


def write_errors(path: Path, errors: list[dict[str, str | int]]) -> None:
    """Write validation errors to CSV."""
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["row", "field", "message"])
        writer.writeheader()
        writer.writerows(errors)
